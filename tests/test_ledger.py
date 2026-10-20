from __future__ import annotations

import zipfile
from datetime import datetime

import pytest
from openpyxl import load_workbook

from movie_agent import config
from movie_agent.ledger import (
    HEADERS,
    RESULT_FOUND,
    RESULT_NOT_FOUND,
    RESULT_NOT_SCREEN,
    Ledger,
    LedgerBusyError,
    LedgerError,
    LedgerRow,
    neutralize_formula,
)

D1 = datetime(2026, 10, 15, 12, 30, 45)
D2 = datetime(2026, 10, 16, 8, 0, 0)


def row(when=D1, uploader="小明", title="全面啟動", result=RESULT_FOUND, text="", images=("screenshots/a.jpg",)):
    return LedgerRow(
        when=when, uploader=uploader, query_type="截圖", query_text=text, result=result,
        title=title, original_title="Inception", year=2010, work_type="電影",
        scene="走廊打鬥", confidence="高", image_paths=list(images),
    )


def not_screen(when=D1):
    return LedgerRow(when, "小明", "截圖", "", RESULT_NOT_SCREEN, scene="一隻貓", image_paths=["screenshots/cat.jpg"])


def sheet(path):
    return load_workbook(path).worksheets[0]


def test_creates_file_with_headers(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    assert ledger.append([row()]) == 1
    ws = sheet(tmp_path / "log.xlsx")
    assert [c.value for c in ws[1]] == HEADERS
    assert [c.value for c in ws[2]] == [
        D1, "小明", "截圖", None, "找到", "全面啟動", "Inception", 2010, "電影", "走廊打鬥", "高", 1,
        "screenshots/a.jpg",
    ]


def test_appends_to_the_end_without_overwriting(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    for t in ("第一部", "第二部", "第三部"):
        ledger.append([row(title=t)])
    ws = sheet(tmp_path / "log.xlsx")
    assert ws.max_row == 4
    assert [ws.cell(r, 6).value for r in (2, 3, 4)] == ["第一部", "第二部", "第三部"]


def test_daily_count_is_whole_channel_not_per_person(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    assert ledger.append([row(uploader="小明")]) == 1
    assert ledger.append([row(uploader="小華")]) == 2  # 不分人
    assert ledger.append([row(uploader="小明")]) == 3
    ws = sheet(tmp_path / "log.xlsx")
    assert [ws.cell(r, 12).value for r in (2, 3, 4)] == [1, 2, 3]


def test_not_found_and_not_screen_do_not_count(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row()])
    assert ledger.append([row(title="", result=RESULT_NOT_FOUND)]) == 1
    assert ledger.append([not_screen()]) == 1
    ws = sheet(tmp_path / "log.xlsx")
    assert [c.value for c in ws[4]][4:12] == ["不是影視畫面", None, None, None, None, "一隻貓", None, 1]  # 空白文字讀回來是 None


def test_daily_count_resets_on_new_day(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row(when=D1)])
    assert ledger.append([row(when=D2)]) == 1


def test_midnight_boundary(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row(when=datetime(2026, 10, 15, 23, 59, 59))])
    assert ledger.append([row(when=datetime(2026, 10, 16, 0, 0, 0))]) == 1


def test_count_follows_hand_edited_result(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row()])
    ledger.append([row()])
    wb = load_workbook(tmp_path / "log.xlsx")
    wb.worksheets[0].cell(2, 5).value = "找不到"  # 你發現第一筆查錯了，手動改掉
    wb.save(tmp_path / "log.xlsx")
    assert ledger.append([row()]) == 2


def test_multiple_rows_from_one_message_share_the_count(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    assert ledger.append([row(), not_screen()]) == 1
    ws = sheet(tmp_path / "log.xlsx")
    assert ws.cell(2, 12).value == 1 and ws.cell(3, 12).value == 1


def test_multiple_images_in_one_cell(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row(images=("screenshots/a.jpg", "screenshots/b.jpg"))])
    assert sheet(tmp_path / "log.xlsx").cell(2, 13).value == "screenshots/a.jpg\nscreenshots/b.jpg"


# ---- 防護 4 ----
@pytest.mark.parametrize("evil", ["=1+1", "+cmd|' /C calc'!A0", "-2", "@SUM(1)", '=HYPERLINK("http://x","點我")'])
def test_formula_text_is_stored_as_text(tmp_path, evil):
    path = tmp_path / "log.xlsx"
    Ledger(path).append([row(uploader=evil, title=evil, text=evil)])
    ws = sheet(path)
    for col in (2, 4, 6):
        cell = ws.cell(2, col)
        assert cell.data_type == "s"
        assert cell.value == "'" + evil
    # xlsx 內部的檔案裡完全沒有公式標籤
    with zipfile.ZipFile(path) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f>" not in xml and "<f " not in xml


def test_neutralize_formula_leaves_normal_text_alone():
    assert neutralize_formula("全面啟動") == "全面啟動"
    assert neutralize_formula("") == ""
    assert neutralize_formula("a=b") == "a=b"
    assert neutralize_formula("\t=1") == "'\t=1"


# ---- 容錯 ----
def test_header_mismatch_refuses_to_write(tmp_path):
    path = tmp_path / "log.xlsx"
    Ledger(path).append([row()])
    wb = load_workbook(path)
    wb.worksheets[0].cell(1, 1).value = "被改過的標題"
    wb.save(path)
    with pytest.raises(LedgerError):
        Ledger(path).append([row()])
    assert sheet(path).max_row == 2  # 沒有被寫入


def test_corrupted_file_raises_ledger_error_and_is_kept(tmp_path):
    path = tmp_path / "log.xlsx"
    path.write_bytes(b"this is not a zip")
    with pytest.raises(LedgerError):
        Ledger(path).append([row()])
    assert path.read_bytes() == b"this is not a zip"


def test_locked_file_warns_retries_and_does_not_crash(tmp_path, monkeypatch, caplog):
    path = tmp_path / "log.xlsx"
    Ledger(path).append([row()])
    monkeypatch.setattr(config, "EXCEL_RETRY_TIMES", 3)
    monkeypatch.setattr(config, "EXCEL_RETRY_WAIT_S", 0)
    calls = []

    def locked(src, dst):
        calls.append(1)
        raise PermissionError("[WinError 32] 另一個程序正在使用此檔案")

    monkeypatch.setattr("movie_agent.ledger.os.replace", locked)
    with pytest.raises(LedgerBusyError):
        Ledger(path).append([row()])
    assert len(calls) == 3
    warnings = [r.getMessage() for r in caplog.records if r.levelname == "WARNING" and "[警告]" in r.getMessage()]
    assert len(warnings) == 3 and "1/3" in warnings[0] and "3/3" in warnings[2]
    assert sheet(path).max_row == 2  # 原檔沒被動到
    assert not list(tmp_path.glob("*.tmp.xlsx"))  # 暫存檔有清掉


def test_succeeds_after_file_is_released(tmp_path, monkeypatch):
    path = tmp_path / "log.xlsx"
    Ledger(path).append([row()])
    monkeypatch.setattr(config, "EXCEL_RETRY_WAIT_S", 0)
    import os

    real_replace = os.replace
    state = {"n": 0}

    def flaky(src, dst):
        state["n"] += 1
        if state["n"] < 3:
            raise PermissionError("locked")
        return real_replace(src, dst)

    monkeypatch.setattr("movie_agent.ledger.os.replace", flaky)
    assert Ledger(path).append([row()]) == 2
    assert sheet(path).max_row == 3
