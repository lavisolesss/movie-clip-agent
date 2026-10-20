from __future__ import annotations

import asyncio
from datetime import datetime

from conftest import found_analysis, found_work, make_job, non_screen_analysis, not_found_work, text_only_analysis
from movie_agent import config
from movie_agent.ledger import Ledger
from movie_agent.pipeline import process_job
from movie_agent.schema import Analysis, AnalysisError, ImageVerdict
from openpyxl import load_workbook


def _run(job, settings, analysis=None, error=None):
    async def fake_analyze(images, text, s):
        if error:
            raise error
        return analysis

    ledger = Ledger(settings.ledger_path)
    return asyncio.run(process_job(job, settings, ledger, analyze_fn=fake_analyze))


def _rows(settings):
    return list(load_workbook(settings.ledger_path).worksheets[0].iter_rows(min_row=2, values_only=True))


def _images(settings):
    return sorted(p.name for p in settings.screenshots_dir.iterdir())


def test_screenshot_flow(settings):
    out = _run(make_job(), settings, found_analysis())
    assert out.ok
    assert _images(settings) == ["2026-10-15 12-30-45 全面啟動.jpg"]
    (row,) = _rows(settings)
    assert row[1:6] == ("小明", "截圖", None, "找到", "全面啟動")
    assert row[6:13] == ("Inception", 2010, "電影", "柯柏在旋轉的走廊裡和保鏢打鬥", "高", 1,
                         "screenshots/2026-10-15 12-30-45 全面啟動.jpg")
    for needle in ("<@1>", "找到了：《全面啟動》", "信心程度：高", "今天頻道裡已找到 1 部作品"):
        assert needle in out.text


def test_text_only_flow_records_the_quote_without_images(settings):
    out = _run(make_job(n_images=0, text="人生就像一盒巧克力"), settings, text_only_analysis())
    assert out.ok
    assert _images(settings) == []
    (row,) = _rows(settings)
    assert row[2:6] == ("文字", "人生就像一盒巧克力", "找到", "阿甘正傳")
    assert row[12] is None  # 沒有截圖
    assert "找到了：《阿甘正傳》" in out.text


def test_screenshot_plus_text(settings):
    _run(make_job(text="You're waiting for a train"), settings, found_analysis())
    (row,) = _rows(settings)
    assert row[2] == "截圖+文字" and row[3] == "You're waiting for a train"


def test_daily_count_accumulates_across_people(settings):
    _run(make_job(uid=1, name="小明"), settings, found_analysis())
    out = _run(make_job(uid=2, name="小華"), settings, found_analysis(title="黑暗騎士"))
    assert "今天頻道裡已找到 2 部作品" in out.text
    assert [r[11] for r in _rows(settings)] == [1, 2]


def test_not_found_is_recorded_but_not_counted(settings):
    out = _run(make_job(), settings, Analysis((ImageVerdict(1, True, ""),), not_found_work()))
    assert out.ok
    assert _images(settings) == ["2026-10-15 12-30-45 未找到.jpg"]
    (row,) = _rows(settings)
    assert row[4] == "找不到" and row[5] is None and row[9] == "畫面太暗，看不出線索" and row[11] == 0
    assert "沒有找到確定的作品" in out.text


def test_non_screen_is_recorded_saved_and_replied(settings):
    out = _run(make_job(), settings, non_screen_analysis("馬克杯"))
    assert out.ok
    assert _images(settings) == ["2026-10-15 12-30-45 不是影視畫面_馬克杯.jpg"]
    (row,) = _rows(settings)
    assert row[4] == "不是影視畫面" and row[9] == "馬克杯" and row[11] == 0
    assert "你傳的不是影視畫面，是馬克杯。" in out.text


def test_mixed_message_makes_two_rows(settings):
    analysis = Analysis(
        images=(ImageVerdict(1, True, ""), ImageVerdict(2, True, ""), ImageVerdict(3, False, "一隻貓")),
        work=found_work(),
    )
    out = _run(make_job(n_images=3), settings, analysis)
    assert out.ok
    rows = _rows(settings)
    assert [r[4] for r in rows] == ["找到", "不是影視畫面"]
    assert rows[0][12].splitlines() == [
        "screenshots/2026-10-15 12-30-45 全面啟動.jpg",
        "screenshots/2026-10-15 12-30-45 全面啟動-2.jpg",
    ]
    assert [r[11] for r in rows] == [1, 1]
    assert "截圖（第 1、2 張）找到了" in out.text and "第 3 張不是影視畫面，是一隻貓。" in out.text
    assert len(_images(settings)) == 3


def test_analysis_failure_leaves_no_trace(settings):
    out = _run(make_job(), settings, error=AnalysisError("boom"))
    assert not out.ok and out.text == "查詢失敗，請再傳一次"
    assert _images(settings) == []
    assert not settings.ledger_path.exists()


def test_unexpected_error_does_not_crash(settings):
    out = _run(make_job(), settings, error=RuntimeError("something odd"))
    assert not out.ok and out.text == "查詢失敗，請再傳一次"
    assert _images(settings) == []


def test_busy_excel_deletes_saved_images_and_replies(settings, monkeypatch):
    monkeypatch.setattr(config, "EXCEL_RETRY_TIMES", 2)
    monkeypatch.setattr(config, "EXCEL_RETRY_WAIT_S", 0)

    def locked(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr("movie_agent.ledger.os.replace", locked)
    out = _run(make_job(), settings, found_analysis())
    assert not out.ok and "記錄檔現在寫不進去" in out.text
    assert _images(settings) == []  # 已存的截圖有刪掉


def test_truncated_note_in_reply(settings):
    out = _run(make_job(truncated=True), settings, found_analysis())
    assert f"一次最多處理 {config.MAX_IMAGES_PER_MESSAGE} 張" in out.text


def test_evil_names_from_claude(settings):
    # 防護 3 + 防護 4：claude 回來的作品名稱是惡意的
    evil = "=1+1/..\\..\\x:y"  # 開頭是 =（防護 4），中間有路徑亂跑的字（防護 3）
    out = _run(make_job(), settings, found_analysis(title=evil))
    assert out.ok
    files = [p for p in settings.data_dir.rglob("*") if p.is_file() and p != settings.ledger_path]
    assert files and all(p.parent == settings.screenshots_dir for p in files)
    ws = load_workbook(settings.ledger_path).worksheets[0]
    assert ws.cell(2, 6).data_type == "s" and ws.cell(2, 6).value.startswith("'")


def test_evil_nickname_and_quote_go_to_excel_as_text(settings):
    evil = '=HYPERLINK("http://x","點我")'
    out = _run(make_job(name=evil, text=evil), settings, found_analysis())
    assert out.ok
    ws = load_workbook(settings.ledger_path).worksheets[0]
    for col in (2, 4):
        assert ws.cell(2, col).data_type == "s" and ws.cell(2, col).value.startswith("'=")


def test_message_time_not_processing_time_decides_the_day(settings):
    _run(make_job(when=datetime(2026, 10, 15, 23, 59, 0)), settings, found_analysis())
    out = _run(make_job(when=datetime(2026, 10, 16, 0, 1, 0)), settings, found_analysis())
    assert "今天頻道裡已找到 1 部作品" in out.text
